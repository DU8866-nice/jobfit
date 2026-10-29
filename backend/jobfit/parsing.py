"""Turn an uploaded resume file into plain text."""

from __future__ import annotations

import io
import re
from pathlib import Path

from .errors import AppError

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")
MIN_TEXT_CHARS = 50


class ResumeParseError(AppError):
    """The file could not be turned into usable text."""


class FileTooLargeError(ResumeParseError):
    pass


def extract_text(filename: str, data: bytes, *, max_bytes: int) -> str:
    if len(data) > max_bytes:
        raise FileTooLargeError("file_too_large", f"The file is larger than {max_bytes / 1_000_000:.0f} MB.")

    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ResumeParseError("unsupported_file_type", f"Unsupported file type. Please upload one of: {supported}.")

    text = clean_text(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            "no_text_in_file",
            "Couldn't read enough text from this file. If it is a scanned PDF, paste the text instead.",
        )
    return text


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        encrypted = reader.is_encrypted
        pages = [] if encrypted else [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises many error types for malformed files
        raise ResumeParseError(
            "unreadable_pdf", "This PDF couldn't be read. Try exporting it again, or paste the text."
        ) from exc
    if encrypted:
        raise ResumeParseError("encrypted_pdf", "This PDF is password-protected. Please upload an unlocked copy.")
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated error types for bad files
        raise ResumeParseError(
            "unreadable_docx", "This Word file couldn't be read. Try saving it again, or paste the text."
        ) from exc

    parts = [p.text for p in document.paragraphs]
    # Many resume templates put content in tables.
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells repeat their text
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
